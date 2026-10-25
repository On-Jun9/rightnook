#!/usr/bin/env python3
"""Generate or verify RightNook's bundled blank documents (development only)."""
import argparse
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import posixpath
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "Sources/Core/Resources/FileTemplates"


def write_package(document, path):
    stream = BytesIO()
    document.save(stream)
    with ZipFile(stream) as source, ZipFile(path, "w", ZIP_DEFLATED, compresslevel=9) as target:
        removed = {name for name in source.namelist()
                   if name.startswith(("docProps/thumbnail.", "ppt/printerSettings/", "customXml/"))}
        for name in sorted(source.namelist()):
            if name in removed:
                continue
            data = source.read(name)
            if name.endswith(".rels"):
                tree = ET.fromstring(data)
                base = "" if name == "_rels/.rels" else str(Path(name).parent.parent)
                for relationship in list(tree):
                    destination = posixpath.normpath(posixpath.join(base, relationship.attrib["Target"])).lstrip("/")
                    if destination in removed:
                        tree.remove(relationship)
                data = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            elif name == "[Content_Types].xml":
                tree = ET.fromstring(data)
                remaining = set(source.namelist()) - removed
                for part in list(tree):
                    if part.get("PartName", "").lstrip("/") in removed or (
                        "Extension" in part.attrib and
                        not any(item.endswith("." + part.attrib["Extension"]) for item in remaining)
                    ):
                        tree.remove(part)
                data = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            elif name == "docProps/app.xml":
                # Do not retain the generator's sample application, printer,
                # thumbnail, or editing-session metadata in a blank document.
                data = (b'<?xml version="1.0" encoding="UTF-8"?>'
                        b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                        b'<Application>RightNook</Application></Properties>')
            if name == "docProps/core.xml":
                for field in [b"created", b"modified"]:
                    data = re.sub(rb"(<dcterms:" + field + rb"[^>]*>)[^<]*(</dcterms:" + field + rb">)",
                                  rb"\g<1>2000-01-01T00:00:00Z\g<2>", data)
            entry = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, data, compresslevel=9)


def generate(directory):
    from docx import Document
    from docx.oxml.ns import qn
    from openpyxl import Workbook
    from openpyxl.packaging.core import DocumentProperties
    from pptx import Presentation
    from pptx.util import Inches

    directory.mkdir(parents=True, exist_ok=True)
    word = Document()
    word.add_paragraph()
    for setting in word.settings.element.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")
    slides = Presentation()
    slides.slide_width, slides.slide_height = Inches(13.333333), Inches(7.5)
    slides.slides.add_slide(slides.slide_layouts[6])
    for document in [word, slides]:
        properties = document.core_properties
        for field in ["author", "last_modified_by", "title", "subject", "comments", "keywords", "category"]:
            setattr(properties, field, "")
        properties.created = properties.modified = datetime(2000, 1, 1)
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.properties = DocumentProperties(creator="", lastModifiedBy="", created=datetime(2000, 1, 1), modified=datetime(2000, 1, 1))
    for document, suffix in [(word, "docx"), (workbook, "xlsx"), (slides, "pptx")]:
        write_package(document, directory / ("blank." + suffix))
    (directory / "blank.rtf").write_bytes(b"{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0\\fnil Helvetica;}}\\viewkind4\\uc1\\pard\\f0\\fs24\\par}\n")


def verify(directory):
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    results = {}
    expected = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml", "pptx": "ppt/presentation.xml"}
    for suffix, main_part in expected.items():
        paths = list(directory.glob("*." + suffix))
        if not paths:
            raise RuntimeError("검증할 파일 없음: " + suffix)
        for path in paths:
            with ZipFile(path) as package:
                assert package.testzip() is None, path
                names = set(package.namelist())
                assert {"[Content_Types].xml", "_rels/.rels", main_part} <= names, path
                assert not any("vbaProject" in name or "/embeddings/" in name for name in names), path
                assert all(name.endswith((".xml", ".rels")) for name in names), (path, "unexpected binary part")
                assert not any(name.startswith("customXml/") for name in names), path
                metadata = ET.fromstring(package.read("docProps/core.xml"))
                private_fields = {"creator", "lastModifiedBy", "title", "subject", "description", "keywords", "category"}
                assert all(not (item.text or "").strip() for item in metadata
                           if item.tag.rsplit("}", 1)[-1] in private_fields), (path, "nonempty metadata")
                for name in names:
                    if name.endswith((".xml", ".rels")):
                        tree = ET.fromstring(package.read(name))
                        if name.endswith(".rels"):
                            base = "" if name == "_rels/.rels" else str(Path(name).parent.parent)
                            for relationship in tree:
                                assert relationship.get("TargetMode") != "External", (path, name)
                                destination = relationship.attrib["Target"]
                                resolved = posixpath.normpath(posixpath.join(base, destination)).lstrip("/")
                                assert resolved in names, (path, name, destination)
            if suffix == "docx":
                document = Document(path)
                assert document.paragraphs and all(not paragraph.text for paragraph in document.paragraphs)
                ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                modes = [item.get(ns + "val") for item in document.settings.element.iter(ns + "compatSetting")
                         if item.get(ns + "name") == "compatibilityMode"]
                assert modes == ["15"], (path, modes)
            elif suffix == "xlsx":
                workbook = load_workbook(path)
                assert workbook.sheetnames == ["Sheet1"] and workbook.active["A1"].value is None
                workbook.close()
            else:
                document = Presentation(path)
                assert len(document.slides) == 1 and len(document.slides[0].shapes) == 0
            results[path.name] = {"bytes": path.stat().st_size, "zipXmlRelationshipsAndReader": "passed"}
    for path in directory.glob("*.rtf"):
        assert path.read_bytes().startswith(b"{\\rtf1") and path.read_bytes().rstrip().endswith(b"}")
        results[path.name] = {"bytes": path.stat().st_size, "rtfHeader": "passed"}
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["generate", "verify"])
    parser.add_argument("directory", nargs="?", type=Path, default=TEMPLATES)
    args = parser.parse_args()
    if args.command == "generate":
        generate(args.directory)
    verify(args.directory)
